# -*- coding: utf-8 -*-
"""
Генерация оригинального Excel файла для веб-приложения
"""
import sqlite3
import os
import shutil
from datetime import datetime
from openpyxl import load_workbook


def generate_original_excel_report():
    """Создает заполненный оригинальный Excel файл для веб-приложения"""
    try:
        print("Начинаю создание отчета по оригинальному шаблону...")
        
        # Получаем данные из базы
        conn = sqlite3.connect('livestock.db')
        conn.row_factory = sqlite3.Row
        
        cursor = conn.execute('''
            SELECT farm_name, category, reason, disposal_date, cow_id 
            FROM cows 
            ORDER BY farm_name, disposal_date
        ''')
        
        data = [dict(row) for row in cursor.fetchall()]
        conn.close()
        
        if not data:
            print("Нет данных в базе")
            return None
        
        # Группируем данные по фермам
        farm_stats = {}
        for cow in data:
            farm = cow['farm_name']
            category = cow['category']
            
            if farm not in farm_stats:
                farm_stats[farm] = {
                    'падёж': 0,
                    'выбраковка': 0,
                    'санитарный': 0,
                    'всего': 0
                }
            
            if category in farm_stats[farm]:
                farm_stats[farm][category] += 1
                farm_stats[farm]['всего'] += 1
        
        print(f"Получено {len(farm_stats)} ферм из базы")
        
        # Путь к оригинальному файлу
        original_path = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '2,3Выбытие скота Сводка.xlsx'))
        
        if not os.path.exists(original_path):
            print(f"Файл не найден: {original_path}")
            return None
        
        # Создаем копию
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        new_filename = f'Выбытие_скота_по_шаблону_{timestamp}.xlsx'
        new_path = os.path.join('reports', new_filename)
        
        # Создаем папку reports если её нет
        os.makedirs('reports', exist_ok=True)
        
        # Копируем файл
        shutil.copy2(original_path, new_path)
        print(f"Создана копия: {new_path}")
        
        # Открываем для редактирования
        wb = load_workbook(new_path)
        ws = wb.active
        
        # Находим колонку "Пало"
        пало_col = None
        for col in range(1, ws.max_column + 1):
            for row in range(1, 10):
                cell_value = ws.cell(row=row, column=col).value
                if cell_value and isinstance(cell_value, str) and 'пало' in cell_value.lower():
                    пало_col = col
                    print(f"Найдена колонка 'Пало': колонка {col}")
                    break
            if пало_col:
                break
        
        if not пало_col:
            пало_col = 29  # По умолчанию
        
        # Находим строки с фермами
        start_row = 5
        last_row = start_row
        for row in range(start_row, ws.max_row + 1):
            if ws.cell(row=row, column=1).value is not None:
                last_row = row
            else:
                break
        
        print(f"Данные ферм: строки {start_row}-{last_row}")
        
        # Заполняем существующие фермы
        фермы_обработаны = 0
        
        for row in range(start_row, last_row + 1):
            farm_name_cell = ws.cell(row=row, column=2)
            
            if not farm_name_cell.value or not isinstance(farm_name_cell.value, str):
                continue
            
            farm_name = farm_name_cell.value.strip()
            
            # Ищем ферму в базе данных
            matching_farm = None
            for db_farm in list(farm_stats.keys()):
                db_clean = db_farm.strip().lower().replace('"', '').replace("'", "")
                file_clean = farm_name.strip().lower().replace('"', '').replace("'", "")
                
                if (db_clean in file_clean or 
                    file_clean in db_clean or
                    db_farm.strip() == farm_name.strip()):
                    matching_farm = db_farm
                    break
            
            if matching_farm:
                stats = farm_stats[matching_farm]
                
                # Заполняем "Пало" как сумма
                пало_value = stats['выбраковка'] + stats['санитарный'] + stats['падёж']
                ws.cell(row=row, column=пало_col).value = пало_value
                
                фермы_обработаны += 1
                del farm_stats[matching_farm]
        
        print(f"Обработано ферм в файле: {фермы_обработаны}")
        
        # Добавляем новые фермы в конец
        if farm_stats:
            print(f"Добавляю {len(farm_stats)} новых ферм")
            
            next_row = last_row + 1
            next_number = ws.cell(row=last_row, column=1).value + 1 if ws.cell(row=last_row, column=1).value else 1
            
            for farm, stats in farm_stats.items():
                ws.cell(row=next_row, column=1).value = next_number
                ws.cell(row=next_row, column=2).value = farm
                пало_value = stats['выбраковка'] + stats['санитарный'] + stats['падёж']
                ws.cell(row=next_row, column=пало_col).value = пало_value
                
                next_row += 1
                next_number += 1
        
        # Сохраняем
        wb.save(new_path)
        wb.close()
        
        print(f"Отчет успешно создан: {new_path}")
        return new_path
        
    except Exception as e:
        print(f"Ошибка при создании отчета: {e}")
        import traceback
        traceback.print_exc()
        return None


if __name__ == '__main__':
    result = generate_original_excel_report()
    if result:
        print(f"✅ Файл создан: {result}")
    else:
        print("❌ Не удалось создать файл")