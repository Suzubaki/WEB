# report_generator.py - Генерация первичных документов АПК Республики Беларусь (Форма 209-АПК, 210-АПК)
import sqlite3
import csv
import io
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime
from database import get_db_connection

def generate_csv_report(start_date, end_date, farm_name=None):
    """Генерация расширенного отчета в формате CSV"""
    conn = get_db_connection()
    cursor = conn.cursor()
    
    where_clause = "WHERE disposal_date BETWEEN ? AND ?"
    params = [start_date, end_date]
    
    if farm_name and farm_name != 'all':
        where_clause += " AND farm_name = ?"
        params.append(farm_name)
        
    cursor.execute(f'''
        SELECT cow_id, ear_tag, farm_name, category, reason, disposal_date, lactation, weight, notes, created_at 
        FROM cows 
        {where_clause}
        ORDER BY disposal_date DESC
    ''', params)
    
    rows = cursor.fetchall()
    conn.close()
    
    output = io.StringIO()
    writer = csv.writer(output, delimiter=';', quotechar='"', quoting=csv.QUOTE_MINIMAL)
    
    # Заголовок
    writer.writerow(['ID записи', 'Идентиф. номер AITS (РБ) / Инв. №', 'Ферма / МТФ', 'Категория', 'Причина (диагноз)', 'Дата выбытия', 'Лактация/Возраст', 'Живая масса (кг)', 'Заключение ветврача', 'Дата внесения'])
    
    for row in rows:
        writer.writerow([
            row['cow_id'],
            row['ear_tag'] or '-',
            row['farm_name'],
            row['category'],
            row['reason'],
            row['disposal_date'],
            row['lactation'] or '-',
            row['weight'] or '-',
            row['notes'] or '-',
            row['created_at']
        ])
        
    output.seek(0)
    return output.getvalue()

def generate_simple_excel_report(start_date, end_date, farm_name=None):
    """Генерация простого Excel отчета со всеми полями"""
    conn = get_db_connection()
    cursor = conn.cursor()
    
    where_clause = "WHERE disposal_date BETWEEN ? AND ?"
    params = [start_date, end_date]
    
    if farm_name and farm_name != 'all':
        where_clause += " AND farm_name = ?"
        params.append(farm_name)
        
    cursor.execute(f'''
        SELECT cow_id, ear_tag, farm_name, category, reason, disposal_date, lactation, weight, notes 
        FROM cows 
        {where_clause}
        ORDER BY farm_name, disposal_date DESC
    ''', params)
    
    rows = cursor.fetchall()
    conn.close()
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Ведомость выбытия"
    
    ws.append(['ID записи', 'Идентиф. номер AITS (РБ) / Бирка', 'Ферма / МТФ', 'Категория', 'Причина выбытия', 'Дата выбытия', 'Лактация', 'Масса (кг)', 'Заключение ветврача'])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        
    for row in rows:
        ws.append([
            row['cow_id'],
            row['ear_tag'] or '-',
            row['farm_name'],
            row['category'],
            row['reason'],
            row['disposal_date'],
            row['lactation'] or '-',
            row['weight'] or '-',
            row['notes'] or '-'
        ])
        
    ws.column_dimensions['A'].width = 16
    ws.column_dimensions['B'].width = 22
    ws.column_dimensions['C'].width = 18
    ws.column_dimensions['D'].width = 16
    ws.column_dimensions['E'].width = 28
    ws.column_dimensions['F'].width = 14
    ws.column_dimensions['G'].width = 12
    ws.column_dimensions['H'].width = 14
    ws.column_dimensions['I'].width = 30
        
    return wb

def generate_form_209_apk_excel(start_date, end_date, farm_name=None):
    """
    Генерация официального Акта на выбытие животных и птицы (забой, прирезка и падёж)
    Типовая форма 209-АПК, утвержденная Министерством сельского хозяйства и продовольствия Республики Беларусь
    """
    conn = get_db_connection()
    cursor = conn.cursor()
    
    where_clause = "WHERE disposal_date BETWEEN ? AND ?"
    params = [start_date, end_date]
    if farm_name and farm_name != 'all':
        where_clause += " AND farm_name = ?"
        params.append(farm_name)
        
    cursor.execute(f'''
        SELECT cow_id, ear_tag, farm_name, category, reason, disposal_date, lactation, weight, notes 
        FROM cows 
        {where_clause}
        ORDER BY disposal_date ASC
    ''', params)
    rows = cursor.fetchall()
    conn.close()
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Форма 209-АПК (РБ)"
    
    # Стили
    title_font = Font(name='Arial', size=11, bold=True)
    header_font = Font(name='Arial', size=9, bold=True)
    body_font = Font(name='Arial', size=9)
    bold_body_font = Font(name='Arial', size=9, bold=True)
    small_font = Font(name='Arial', size=8, italic=True)
    
    thin_border = Border(
        left=Side(style='thin', color='000000'),
        right=Side(style='thin', color='000000'),
        top=Side(style='thin', color='000000'),
        bottom=Side(style='thin', color='000000')
    )
    header_fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
    
    # Шапка формы 209-АПК (РБ)
    ws.merge_cells('F1:I1')
    ws['F1'] = "Типовая форма 209-АПК"
    ws['F1'].font = Font(name='Arial', size=9, bold=True)
    ws['F1'].alignment = Alignment(horizontal='right')
    
    ws.merge_cells('E2:I2')
    ws['E2'] = "Утверждена постановлением Минсельхозпрода Республики Беларусь"
    ws['E2'].font = small_font
    ws['E2'].alignment = Alignment(horizontal='right')
    
    ws['A4'] = f"Сельскохозяйственная организация: {farm_name if (farm_name and farm_name != 'all') else 'Сводный акт по хозяйству (фермам)'}"
    ws['A4'].font = Font(name='Arial', size=10, bold=True)
    
    ws.merge_cells('A6:I6')
    ws['A6'] = "АКТ НА ВЫБЫТИЕ ЖИВОТНЫХ И ПТИЦЫ (ЗАБОЙ, ПРИРЕЗКА И ПАДЁЖ)"
    ws['A6'].font = title_font
    ws['A6'].alignment = Alignment(horizontal='center', vertical='center')
    
    ws.merge_cells('A7:I7')
    ws['A7'] = f"за период с {start_date} по {end_date} года"
    ws['A7'].font = body_font
    ws['A7'].alignment = Alignment(horizontal='center', vertical='center')
    
    # Табличная часть
    headers = [
        ('№ п/п', 5),
        ('Учетный номер', 14),
        ('Идент. № (AITS / Бирка)', 18),
        ('Подразделение / МТФ', 18),
        ('Вид выбытия (РБ)', 16),
        ('Причина выбытия (диагноз)', 25),
        ('Лактация / Возраст', 15),
        ('Живая масса (кг)', 14),
        ('Дата выбытия', 13)
    ]
    
    row_num = 9
    for col_idx, (h_title, width) in enumerate(headers, 1):
        cell = ws.cell(row=row_num, column=col_idx, value=h_title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = thin_border
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = width
    ws.row_dimensions[row_num].height = 28
    
    total_weight = 0.0
    align_center = Alignment(horizontal='center', vertical='center')
    align_left = Alignment(horizontal='left', vertical='center')

    for idx, r in enumerate(rows, 1):
        row_num += 1
        ws.row_dimensions[row_num].height = 20
        
        w_val = r['weight'] or 0.0
        total_weight += float(w_val)
        
        vals = [
            idx,
            r['cow_id'],
            r['ear_tag'] or '-',
            r['farm_name'],
            r['category'].capitalize(),
            r['reason'],
            r['lactation'] or '-',
            r['weight'] if r['weight'] is not None else '-',
            r['disposal_date']
        ]
        
        for col_idx, val in enumerate(vals, 1):
            cell = ws.cell(row=row_num, column=col_idx, value=val)
            cell.font = body_font
            cell.border = thin_border
            if col_idx in [1, 2, 3, 7, 8, 9]:
                cell.alignment = align_center
            else:
                cell.alignment = align_left
                
    # Строка Итого
    row_num += 1
    ws.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=7)
    ws.cell(row=row_num, column=1, value="ИТОГО ВЫБЫЛО ПО АКТУ:").font = bold_body_font
    ws.cell(row=row_num, column=1).alignment = Alignment(horizontal='right', vertical='center')
    for c in range(1, 8):
        ws.cell(row=row_num, column=c).border = thin_border
        ws.cell(row=row_num, column=c).fill = header_fill
        
    cell_head_cnt = ws.cell(row=row_num, column=8, value=f"{total_weight:.1f} кг" if total_weight > 0 else "-")
    cell_head_cnt.font = bold_body_font
    cell_head_cnt.alignment = Alignment(horizontal='center', vertical='center')
    cell_head_cnt.border = thin_border
    cell_head_cnt.fill = header_fill
    
    cell_total_cows = ws.cell(row=row_num, column=9, value=f"{len(rows)} гол.")
    cell_total_cows.font = bold_body_font
    cell_total_cows.alignment = Alignment(horizontal='center', vertical='center')
    cell_total_cows.border = thin_border
    cell_total_cows.fill = header_fill
    
    # Блок подписей комиссии сельхозорганизации РБ
    row_num += 3
    ws.cell(row=row_num, column=1, value="Комиссия сельхозорганизации:").font = bold_body_font
    row_num += 1
    ws.cell(row=row_num, column=1, value="Руководитель хозяйства (Директор): _______________________ (подпись)").font = body_font
    row_num += 1
    ws.cell(row=row_num, column=1, value="Главный ветеринарный врач:        _______________________ (подпись)").font = body_font
    row_num += 1
    ws.cell(row=row_num, column=1, value="Главный зоотехник:                 _______________________ (подпись)").font = body_font
    row_num += 1
    ws.cell(row=row_num, column=1, value="Заведующий МТФ / фермой:           _______________________ (подпись)").font = body_font
    row_num += 1
    ws.cell(row=row_num, column=1, value="Материально ответственное лицо:    _______________________ (подпись)").font = body_font
    
    return wb

# Синоним для совместимости
generate_sp54_act_excel = generate_form_209_apk_excel
