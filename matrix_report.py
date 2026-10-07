# matrix_report.py - Генерация матричного отчета "фермы х причины"
import sqlite3
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime
from database import get_db_connection, get_reasons_for_category

def generate_matrix_excel_report(category, start_date=None, end_date=None):
    """
    Генерирует Excel отчет в виде матрицы:
    Строки - причины выбытия для указанной категории (из справочника БД)
    Столбцы - фермы
    Значения - количество коров
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Матрица {category.capitalize()}"
    
    # Стили
    title_font = Font(name='Arial', size=13, bold=True, color="1F2937")
    header_font = Font(name='Arial', size=10, bold=True, color="FFFFFF")
    bold_font = Font(name='Arial', size=9, bold=True)
    regular_font = Font(name='Arial', size=9)
    
    header_fill = PatternFill(start_color="3B82F6", end_color="3B82F6", fill_type="solid")
    total_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    
    thin_border = Border(
        left=Side(style='thin', color='D1D5DB'),
        right=Side(style='thin', color='D1D5DB'),
        top=Side(style='thin', color='D1D5DB'),
        bottom=Side(style='thin', color='D1D5DB')
    )
    
    # 1. Получаем список причин для категории из базы данных
    reasons = get_reasons_for_category(category)
    
    # 2. Получаем список всех ферм из БД
    conn = get_db_connection()
    cursor = conn.cursor()
    
    cursor.execute('SELECT DISTINCT farm_name FROM cows WHERE farm_name IS NOT NULL ORDER BY farm_name')
    farms = [row['farm_name'] for row in cursor.fetchall()]
    
    if not farms:
        cursor.execute('SELECT DISTINCT farm_name FROM users WHERE farm_name IS NOT NULL AND farm_name != "" ORDER BY farm_name')
        farms = [row['farm_name'] for row in cursor.fetchall()]
    
    if not farms:
        farms = ['Ферма 1']
    
    # Заголовок
    ws.merge_cells('A1:G1')
    ws['A1'] = "СВОДНАЯ МАТРИЦА ВЫБЫТИЯ СКОТА"
    ws['A1'].font = title_font
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
    
    ws.merge_cells('A2:G2')
    ws['A2'] = f"Категория: {category.upper()}" + (f" | Период: {start_date} — {end_date}" if start_date and end_date else "")
    ws['A2'].font = Font(name='Arial', size=10, italic=True)
    ws['A2'].alignment = Alignment(horizontal='center', vertical='center')
    
    ws.merge_cells('A3:G3')
    ws['A3'] = f"Сформировано: {datetime.now().strftime('%d.%m.%Y %H:%M')}"
    ws['A3'].font = Font(name='Arial', size=9, color="6B7280")
    ws['A3'].alignment = Alignment(horizontal='center', vertical='center')
    
    # Шапка таблицы (строка 5)
    row_num = 5
    ws.cell(row=row_num, column=1, value="ПРИЧИНЫ ВЫБЫТИЯ").font = header_font
    ws.cell(row=row_num, column=1).fill = header_fill
    ws.cell(row=row_num, column=1).alignment = Alignment(horizontal='center', vertical='center')
    ws.cell(row=row_num, column=1).border = thin_border
    
    col_num = 2
    farm_col_map = {}
    for farm in farms:
        cell = ws.cell(row=row_num, column=col_num, value=farm)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = thin_border
        farm_col_map[farm] = col_num
        col_num += 1
        
    # Колонка ИТОГО
    total_col = col_num
    cell = ws.cell(row=row_num, column=total_col, value="ИТОГО")
    cell.font = header_font
    cell.fill = PatternFill(start_color="EF4444", end_color="EF4444", fill_type="solid")
    cell.alignment = Alignment(horizontal='center', vertical='center')
    cell.border = thin_border
    
    # 3. Единый оптимизированный запрос для агрегации данных матрицы (вместо N*M запросов)
    date_filter = ""
    sql_params = [category]
    if start_date and end_date:
        date_filter = "AND disposal_date BETWEEN ? AND ?"
        sql_params.extend([start_date, end_date])

    cursor.execute(f'''
        SELECT reason, farm_name, COUNT(*) as cnt
        FROM cows
        WHERE category = ? {date_filter}
        GROUP BY reason, farm_name
    ''', sql_params)
    
    matrix_counts = {}
    for row in cursor.fetchall():
        matrix_counts[(row['reason'], row['farm_name'])] = row['cnt']
        
    conn.close()

    # Pre-allocate reusable style and alignment objects
    align_center = Alignment(horizontal='center', vertical='center')
    align_left = Alignment(horizontal='left', vertical='center')
    total_col_fill = PatternFill(start_color="EF4444", end_color="EF4444", fill_type="solid")

    farm_totals = {farm: 0 for farm in farms}
    grand_total = 0

    for reason in reasons:
        cell_reason = ws.cell(row=current_row, column=1, value=reason)
        cell_reason.font = regular_font
        cell_reason.alignment = align_left
        cell_reason.border = thin_border
        
        row_total = 0
        for farm in farms:
            count = matrix_counts.get((reason, farm), 0)
            col = farm_col_map[farm]
            
            cell = ws.cell(row=current_row, column=col, value=count if count > 0 else "-")
            cell.font = regular_font
            cell.alignment = align_center
            cell.border = thin_border
            
            row_total += count
            farm_totals[farm] += count
            
        cell = ws.cell(row=current_row, column=total_col, value=row_total)
        cell.font = bold_font
        cell.alignment = align_center
        cell.border = thin_border
        
        grand_total += row_total
        current_row += 1
        
    # Итоговая строка по фермам
    ws.cell(row=current_row, column=1, value="ИТОГО ПО ФЕРМАМ:").font = bold_font
    ws.cell(row=current_row, column=1).fill = total_fill
    ws.cell(row=current_row, column=1).border = thin_border
    
    for farm in farms:
        col = farm_col_map[farm]
        farm_total = farm_totals[farm]
        cell = ws.cell(row=current_row, column=col, value=farm_total)
        cell.font = bold_font
        cell.fill = total_fill
        cell.alignment = align_center
        cell.border = thin_border
        
    cell = ws.cell(row=current_row, column=total_col, value=grand_total)
    cell.font = bold_font
    cell.fill = total_fill
    cell.alignment = align_center
    cell.border = thin_border
    
    ws.column_dimensions['A'].width = 38
    for col in range(2, total_col + 1):
        col_letter = get_column_letter(col)
        ws.column_dimensions[col_letter].width = 16
        
    return wb
