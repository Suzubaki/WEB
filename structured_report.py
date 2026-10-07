# structured_report.py - Генерация структурированного отчета Excel
import sqlite3
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime
from database import get_db_connection

def generate_structured_excel_report(start_date, end_date, farm_name=None):
    """
    Генерирует профессионально оформленный структурированный Excel-отчет
    """
    conn = get_db_connection()
    cursor = conn.cursor()
    
    where_clause = "WHERE disposal_date BETWEEN ? AND ?"
    params = [start_date, end_date]
    
    if farm_name and farm_name != 'all':
        where_clause += " AND farm_name = ?"
        params.append(farm_name)
        
    cursor.execute(f'''
        SELECT cow_id, farm_name, category, reason, disposal_date 
        FROM cows 
        {where_clause}
        ORDER BY farm_name, disposal_date DESC
    ''', params)
    
    rows = cursor.fetchall()
    conn.close()
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Сводный отчёт"
    
    # Стили
    title_font = Font(name='Arial', size=14, bold=True, color="1F497D")
    sub_font = Font(name='Arial', size=10, italic=True)
    header_font = Font(name='Arial', size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    alt_fill = PatternFill(start_color="F2F5F8", end_color="F2F5F8", fill_type="solid")
    
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    
    # Заголовок
    ws.merge_cells('A1:E1')
    ws['A1'] = "СВОДНЫЙ ОТЧЁТ ПО ВЫБЫТИЮ СКОТА"
    ws['A1'].font = title_font
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
    
    ws.merge_cells('A2:E2')
    ws['A2'] = f"Период: с {start_date} по {end_date} | Сформирован: {datetime.now().strftime('%d.%m.%Y %H:%M')}"
    ws['A2'].font = sub_font
    ws['A2'].alignment = Alignment(horizontal='center', vertical='center')
    
    # Шапка
    headers = ['ID коровы', 'Ферма', 'Категория', 'Причина выбытия', 'Дата выбытия']
    ws.row_dimensions[4].height = 25
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col_num, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = thin_border
        
    align_center = Alignment(horizontal='center', vertical='center')
    align_left = Alignment(horizontal='left', vertical='center')

    for r_idx, row in enumerate(rows, 5):
        ws.row_dimensions[r_idx].height = 20
        c1 = ws.cell(row=r_idx, column=1, value=row['cow_id'])
        c2 = ws.cell(row=r_idx, column=2, value=row['farm_name'])
        c3 = ws.cell(row=r_idx, column=3, value=row['category'])
        c4 = ws.cell(row=r_idx, column=4, value=row['reason'])
        c5 = ws.cell(row=r_idx, column=5, value=row['disposal_date'])
        
        c1.alignment = align_center
        c2.alignment = align_left
        c3.alignment = align_center
        c4.alignment = align_left
        c5.alignment = align_center
        
        for c in [c1, c2, c3, c4, c5]:
            c.border = thin_border
            if r_idx % 2 == 0:
                c.fill = alt_fill
                
    ws.column_dimensions['A'].width = 18
    ws.column_dimensions['B'].width = 20
    ws.column_dimensions['C'].width = 18
    ws.column_dimensions['D'].width = 35
    ws.column_dimensions['E'].width = 16
    
    return wb
