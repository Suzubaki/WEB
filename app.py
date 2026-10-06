# app.py - Главный файл приложения Flask
import os
import io
import csv
import json
from datetime import datetime, timedelta
from flask import Flask, render_template, request, redirect, url_for, flash, session, jsonify, send_file, Response
from werkzeug.security import generate_password_hash
import openpyxl

from config import Config
from database import (
    init_db, get_db_connection, generate_cow_id, 
    log_audit_event, get_reasons_for_category
)
from models import User, Cow
from auth import (
    authenticate_user, register_user, get_user_by_id, 
    get_all_users, update_user, delete_user
)
from charts import (
    get_statistics_data, prepare_chart_data, 
    get_dashboard_stats, get_user_statistics, get_anomaly_alerts,
    get_dashboard_full_data
)
from matrix_report import generate_matrix_excel_report
from report_generator import generate_csv_report, generate_simple_excel_report, generate_form_209_apk_excel, generate_sp54_act_excel
from form_210_apk import generate_form_210_apk_excel
from aits_export import generate_aits_csv_registry, generate_aits_xml_registry
from autopsy import get_cow_for_autopsy, save_autopsy_protocol, parse_autopsy_protocol
from economics import calculate_economic_losses
from backup import create_db_backup_bytes, restore_db_from_upload
from database import get_setting, set_setting, get_all_settings
from structured_report import generate_structured_excel_report
from fill_original_excel import generate_official_template_excel

app = Flask(__name__)
app.config.from_object(Config)

# Векторный логотип ОАО «Новая Припять» (Официальный колос)
LOGO_SVG = """<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 240 240" width="100%" height="100%">
  <defs>
    <linearGradient id="wheatGrad1" x1="0%" y1="0%" x2="100%" y2="100%">
      <stop offset="0%" stop-color="#fbbf24" />
      <stop offset="45%" stop-color="#f59e0b" />
      <stop offset="100%" stop-color="#d97706" />
    </linearGradient>
    <linearGradient id="wheatGrad2" x1="100%" y1="0%" x2="0%" y2="100%">
      <stop offset="0%" stop-color="#fde047" />
      <stop offset="50%" stop-color="#eab308" />
      <stop offset="100%" stop-color="#b45309" />
    </linearGradient>
    <linearGradient id="wheatTip" x1="0%" y1="0%" x2="100%" y2="0%">
      <stop offset="0%" stop-color="#f59e0b" />
      <stop offset="100%" stop-color="#f97316" />
    </linearGradient>
    <filter id="softShadow" x="-10%" y="-10%" width="120%" height="120%">
      <feDropShadow dx="0" dy="1.5" stdDeviation="1.5" flood-color="#b45309" flood-opacity="0.25" />
    </filter>
  </defs>
  <rect width="240" height="240" fill="#ffffff" rx="16" />
  <g id="wheat-ear" filter="url(#softShadow)">
    <path d="M42 120 C38 106, 52 92, 60 90 C66 94, 62 108, 48 122 Z" fill="url(#wheatGrad1)" />
    <path d="M44 118 Q54 100 60 92" stroke="#fef08a" stroke-width="1" fill="none" opacity="0.6" />
    <path d="M52 122 C56 112, 70 106, 78 108 C80 116, 68 126, 54 124 Z" fill="url(#wheatGrad2)" />
    <path d="M60 92 C62 80, 80 72, 90 72 C94 78, 88 92, 74 102 Z" fill="url(#wheatGrad1)" />
    <path d="M64 90 Q80 76 89 74" stroke="#fef08a" stroke-width="1.2" fill="none" opacity="0.6" />
    <path d="M76 104 C82 94, 98 90, 106 94 C108 102, 94 114, 80 110 Z" fill="url(#wheatGrad2)" />
    <path d="M90 73 C96 66, 116 64, 126 68 C128 76, 116 88, 102 92 Z" fill="url(#wheatGrad1)" />
    <path d="M96 72 Q114 66 124 70" stroke="#fef08a" stroke-width="1.2" fill="none" opacity="0.6" />
    <path d="M104 94 C112 86, 130 84, 138 90 C138 98, 124 106, 110 102 Z" fill="url(#wheatGrad2)" />
    <path d="M124 70 C134 68, 152 72, 160 80 C158 88, 142 94, 130 90 Z" fill="url(#wheatGrad1)" />
    <path d="M128 72 Q146 72 157 80" stroke="#fef08a" stroke-width="1.2" fill="none" opacity="0.6" />
    <path d="M136 90 C146 86, 160 88, 168 96 C164 104, 150 106, 138 98 Z" fill="url(#wheatGrad2)" />
    <path d="M156 82 C166 82, 180 88, 186 96 C182 102, 168 104, 158 96 Z" fill="url(#wheatGrad1)" />
    <path d="M168 96 C176 96, 192 100, 196 106 C190 110, 178 110, 170 102 Z" fill="url(#wheatTip)" />
    <path d="M184 92 C196 90, 206 91, 210 93" stroke="#eab308" stroke-width="1.5" stroke-linecap="round" fill="none" />
    <path d="M190 98 C202 96, 212 97, 216 100" stroke="#f59e0b" stroke-width="1.8" stroke-linecap="round" fill="none" />
    <path d="M194 104 C204 102, 214 105, 218 108" stroke="#d97706" stroke-width="1.6" stroke-linecap="round" fill="none" />
    <path d="M188 108 C198 108, 208 112, 212 116" stroke="#b45309" stroke-width="1.3" stroke-linecap="round" fill="none" />
    <path d="M174 104 C184 106, 196 112, 202 118" stroke="#d97706" stroke-width="1.2" stroke-linecap="round" fill="none" />
  </g>
  <text x="120" y="148" font-family="'Trebuchet MS', 'Segoe UI', Arial, sans-serif" font-size="25.5" font-weight="900" font-style="italic" fill="#212529" text-anchor="middle" letter-spacing="-0.3">Новая Припять</text>
  <text x="120" y="166" font-family="'Segoe UI', Roboto, Helvetica, Arial, sans-serif" font-size="8.2" font-weight="700" fill="#262a2e" text-anchor="middle" letter-spacing="1.2">ОТКРЫТОЕ АКЦИОНЕРНОЕ ОБЩЕСТВО</text>
</svg>"""

def _serve_logo_file(filename):
    # Поиск файла в локальных папках
    search_dirs = [
        os.path.join(app.root_path, 'static', 'img'),
        os.path.join(app.root_path, 'public', 'img'),
        os.path.join(app.root_path, 'img'),
        os.path.join(app.root_path, 'static')
    ]
    for d in search_dirs:
        candidate = os.path.join(d, filename)
        if os.path.exists(candidate) and os.path.isfile(candidate):
            mimetype = 'image/jpeg' if filename.endswith('.jpg') or filename.endswith('.jpeg') else ('image/png' if filename.endswith('.png') else 'image/svg+xml')
            return send_file(candidate, mimetype=mimetype)
            
    # Если запрашивается jpg, но есть svg (или нет на диске) - возвращаем SVG
    return Response(LOGO_SVG, mimetype='image/svg+xml')

@app.route('/img/<path:filename>')
@app.route('/static/img/<path:filename>')
def serve_img_routes(filename):
    return _serve_logo_file(filename)

@app.route('/favicon.ico')
@app.route('/favicon.png')
@app.route('/favicon.svg')
def serve_favicon():
    return Response(LOGO_SVG, mimetype='image/svg+xml')

@app.route('/logo.jpg')
@app.route('/logo.svg')
@app.route('/img/logo.jpg')
@app.route('/img/logo.svg')
def serve_logo_direct():
    return _serve_logo_file('logo.svg')

def login_required(f):
    from functools import wraps
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if 'user_id' not in session:
            flash('Пожалуйста, войдите в систему', 'warning')
            return redirect(url_for('login'))
        return f(*args, **kwargs)
    return decorated_function

def admin_required(f):
    from functools import wraps
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if 'user_id' not in session:
            flash('Пожалуйста, войдите в систему', 'warning')
            return redirect(url_for('login'))
        if session.get('user_type') != 'admin':
            flash('Доступ запрещен. Требуются права администратора', 'danger')
            return redirect(url_for('dashboard'))
        return f(*args, **kwargs)
    return decorated_function

@app.before_request
def sync_current_user_session():
    if 'user_id' in session:
        user = get_user_by_id(session['user_id'])
        if user:
            session['farm_name'] = user.farm_name
            session['username'] = user.username
            session['user_type'] = user.user_type

# -------------------- ОСНОВНЫЕ МАРШРУТЫ --------------------

@app.route('/')
def index():
    if 'user_id' in session:
        return redirect(url_for('dashboard'))
    return render_template('index.html')

@app.route('/login', methods=['GET', 'POST'])
def login():
    if 'user_id' in session:
        return redirect(url_for('dashboard'))
        
    if request.method == 'POST':
        username = request.form.get('username', '').strip()
        password = request.form.get('password', '')
        
        user = authenticate_user(username, password)
        if user:
            session['user_id'] = user.id
            session['username'] = user.username
            session['user_type'] = user.user_type
            session['farm_name'] = user.farm_name
            
            log_audit_event(
                username=user.username,
                farm_name=user.farm_name,
                action='LOGIN',
                entity_type='USER',
                entity_id=user.id,
                details='Успешная авторизация в системе',
                user_id=user.id,
                ip_address=request.remote_addr
            )
            
            flash(f'Добро пожаловать, {user.username}!', 'success')
            return redirect(url_for('dashboard'))
        else:
            flash('Неверное имя пользователя или пароль', 'danger')
            
    return render_template('login.html')

@app.route('/logout')
def logout():
    username = session.get('username', 'Unknown')
    user_id = session.get('user_id')
    farm_name = session.get('farm_name')
    
    if user_id:
        log_audit_event(
            username=username,
            farm_name=farm_name,
            action='LOGOUT',
            entity_type='USER',
            entity_id=user_id,
            details='Выход из системы',
            user_id=user_id,
            ip_address=request.remote_addr
        )
        
    session.clear()
    flash('Вы вышли из системы', 'info')
    return redirect(url_for('login'))

@app.route('/dashboard')
@login_required
def dashboard():
    user_type = session.get('user_type')
    farm_name = session.get('farm_name')
    
    period = request.args.get('period', 'all')
    custom_start = request.args.get('custom_start')
    custom_end = request.args.get('custom_end')
    
    selected_farm = farm_name
    if user_type == 'admin':
        selected_farm = request.args.get('farm', 'all')
        
    dashboard_data = get_dashboard_full_data(
        user_type=user_type,
        farm_name=selected_farm,
        period=period,
        custom_start=custom_start,
        custom_end=custom_end
    )
    alerts = get_anomaly_alerts(user_type, selected_farm)
    user_stats_data = get_user_statistics() if user_type == 'admin' else None
    
    return render_template('dashboard.html',
                           user_type=user_type,
                           farm_name=farm_name,
                           selected_farm=selected_farm,
                           d=dashboard_data,
                           alerts=alerts,
                           user_stats_data=user_stats_data)

# -------------------- УПРАВЛЕНИЕ ЗАПИСЯМИ О ВЫБЫТИИ --------------------

@app.route('/cows')
@login_required
def view_cows():
    user_type = session.get('user_type')
    farm_name = session.get('farm_name')
    
    conn = get_db_connection()
    cursor = conn.cursor()
    
    if user_type == 'admin':
        cursor.execute('SELECT * FROM cows ORDER BY disposal_date DESC, id DESC')
    else:
        cursor.execute('SELECT * FROM cows WHERE farm_name = ? ORDER BY disposal_date DESC, id DESC', (farm_name,))
        
    cows_data = cursor.fetchall()
    conn.close()
    
    cows = [dict(c) for c in cows_data]
    return render_template('view_cows.html', cows=cows, user_type=user_type)

@app.route('/add_cow', methods=['GET', 'POST'])
@login_required
def add_cow():
    if session.get('user_type') == 'admin':
        flash('Администраторы не могут добавлять коров', 'warning')
        return redirect(url_for('dashboard'))
        
    farm_name = session.get('farm_name')
    
    if request.method == 'POST':
        category = request.form.get('category')
        reason = request.form.get('reason')
        disposal_date = request.form.get('disposal_date')
        ear_tag = request.form.get('ear_tag', '').strip() or None
        lactation = request.form.get('lactation') or None
        weight = request.form.get('weight') or None
        notes = request.form.get('notes', '').strip() or None
        age_group = request.form.get('age_group') or 'Коровы дойного стада'
        breed = request.form.get('breed') or 'Черно-пёстрая'
        milk_yield = request.form.get('milk_yield') or None
        book_value = request.form.get('book_value') or None
        
        if not category or not reason or not disposal_date:
            flash('Заполните обязательные поля: категория, причина и дата', 'danger')
            return render_template('add_cow.html', categories=Config.DISPOSAL_CATEGORIES, farm_name=farm_name, breeds=Config.BREEDS, age_groups=Config.AGE_GROUPS)
            
        today = datetime.now().strftime('%Y-%m-%d')
        if disposal_date > today:
            flash('Дата выбытия не может быть позже сегодняшнего дня', 'danger')
            return render_template('add_cow.html', categories=Config.DISPOSAL_CATEGORIES, farm_name=farm_name, breeds=Config.BREEDS, age_groups=Config.AGE_GROUPS)
            
        cow_id = generate_cow_id(farm_name)
        
        conn = get_db_connection()
        cursor = conn.cursor()
        cursor.execute('''
            INSERT INTO cows (cow_id, ear_tag, farm_name, category, reason, disposal_date, 
                              lactation, weight, notes, age_group, breed, milk_yield, book_value, created_by)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        ''', (cow_id, ear_tag, farm_name, category, reason, disposal_date, 
              lactation, weight, notes, age_group, breed, milk_yield, book_value, session.get('user_id')))
        new_id = cursor.lastrowid
        conn.commit()
        conn.close()
        
        log_audit_event(
            username=session.get('username'),
            farm_name=farm_name,
            action='CREATE',
            entity_type='COW',
            entity_id=new_id,
            details=f"Добавлена корова {cow_id} (Ушной № {ear_tag or '-'}), {category}: {reason}",
            user_id=session.get('user_id'),
            ip_address=request.remote_addr
        )
        
        flash(f'Корова успешно добавлена с системным номером: {cow_id}', 'success')
        return redirect(url_for('view_cows'))
        
    return render_template('add_cow.html', categories=Config.DISPOSAL_CATEGORIES, farm_name=farm_name, breeds=Config.BREEDS, age_groups=Config.AGE_GROUPS)

@app.route('/edit_cow/<int:cow_id>', methods=['GET', 'POST'])
@login_required
def edit_cow(cow_id):
    user_type = session.get('user_type')
    farm_name = session.get('farm_name')
    
    conn = get_db_connection()
    cursor = conn.cursor()
    
    if user_type == 'admin':
        cursor.execute('SELECT * FROM cows WHERE id = ?', (cow_id,))
    else:
        cursor.execute('SELECT * FROM cows WHERE id = ? AND farm_name = ?', (cow_id, farm_name))
        
    cow_row = cursor.fetchone()
    
    if not cow_row:
        conn.close()
        flash('Запись не найдена или у вас нет прав на её редактирование', 'danger')
        return redirect(url_for('view_cows'))
        
    cow = dict(cow_row)
    
    if request.method == 'POST':
        category = request.form.get('category')
        reason = request.form.get('reason')
        disposal_date = request.form.get('disposal_date')
        ear_tag = request.form.get('ear_tag', '').strip() or None
        lactation = request.form.get('lactation') or None
        weight = request.form.get('weight') or None
        notes = request.form.get('notes', '').strip() or None
        age_group = request.form.get('age_group') or 'Коровы дойного стада'
        breed = request.form.get('breed') or 'Черно-пёстрая'
        milk_yield = request.form.get('milk_yield') or None
        book_value = request.form.get('book_value') or None
        
        today = datetime.now().strftime('%Y-%m-%d')
        if disposal_date > today:
            flash('Дата выбытия не может быть в будущем', 'danger')
            conn.close()
            return render_template('edit_cow.html', cow=cow, categories=Config.DISPOSAL_CATEGORIES)
            
        cursor.execute('''
            UPDATE cows SET 
                category = ?, 
                reason = ?, 
                disposal_date = ?, 
                ear_tag = ?, 
                lactation = ?, 
                weight = ?, 
                notes = ?, 
                age_group = ?,
                breed = ?,
                milk_yield = ?,
                book_value = ?,
                updated_at = CURRENT_TIMESTAMP, 
                updated_by = ?
            WHERE id = ?
        ''', (category, reason, disposal_date, ear_tag, lactation, weight, notes, 
              age_group, breed, milk_yield, book_value, session.get('user_id'), cow_id))
        conn.commit()
        conn.close()
        
        log_audit_event(
            username=session.get('username'),
            farm_name=cow['farm_name'],
            action='UPDATE',
            entity_type='COW',
            entity_id=cow_id,
            details=f"Отредактирована запись {cow['cow_id']}. Новые данные: {category} - {reason}, дата {disposal_date}",
            user_id=session.get('user_id'),
            ip_address=request.remote_addr
        )
        
        flash(f'Запись {cow["cow_id"]} успешно обновлена', 'success')
        return redirect(url_for('view_cows'))
        
    conn.close()
    return render_template('edit_cow.html', cow=cow, categories=Config.DISPOSAL_CATEGORIES)

@app.route('/delete_cow/<int:cow_id>', methods=['POST'])
@login_required
def delete_cow(cow_id):
    user_type = session.get('user_type')
    farm_name = session.get('farm_name')
    
    conn = get_db_connection()
    cursor = conn.cursor()
    
    if user_type == 'admin':
        cursor.execute('SELECT * FROM cows WHERE id = ?', (cow_id,))
    else:
        cursor.execute('SELECT * FROM cows WHERE id = ? AND farm_name = ?', (cow_id, farm_name))
        
    cow = cursor.fetchone()
    if not cow:
        conn.close()
        if request.is_json:
            return jsonify({'error': 'Запись не найдена'}), 404
        flash('Запись не найдена', 'danger')
        return redirect(url_for('view_cows'))
        
    cow_dict = dict(cow)
    cursor.execute('DELETE FROM cows WHERE id = ?', (cow_id,))
    conn.commit()
    conn.close()
    
    log_audit_event(
        username=session.get('username'),
        farm_name=cow_dict['farm_name'],
        action='DELETE',
        entity_type='COW',
        entity_id=cow_id,
        details=f"Удалена запись {cow_dict['cow_id']} ({cow_dict['category']}: {cow_dict['reason']})",
        user_id=session.get('user_id'),
        ip_address=request.remote_addr
    )
    
    if request.is_json or request.headers.get('X-Requested-With') == 'XMLHttpRequest':
        return jsonify({'success': True, 'message': 'Запись удалена'})
        
    flash(f'Запись {cow_dict["cow_id"]} успешно удалена', 'success')
    return redirect(url_for('view_cows'))

@app.route('/bulk_delete_cows', methods=['POST'])
@app.route('/delete_multiple_cows', methods=['POST'], endpoint='delete_multiple_cows')
@login_required
def bulk_delete_cows():
    """Массовое удаление записей"""
    data = request.get_json() or {}
    cow_ids = data.get('ids', []) or data.get('cow_ids', [])
    if not cow_ids:
        return jsonify({'error': 'Не выбрано ни одной записи'}), 400
        
    user_type = session.get('user_type')
    farm_name = session.get('farm_name')
    
    conn = get_db_connection()
    cursor = conn.cursor()
    
    placeholders = ','.join('?' * len(cow_ids))
    if user_type == 'admin':
        query = f"SELECT id, cow_id, farm_name FROM cows WHERE id IN ({placeholders})"
        cursor.execute(query, cow_ids)
    else:
        query = f"SELECT id, cow_id, farm_name FROM cows WHERE id IN ({placeholders}) AND farm_name = ?"
        cursor.execute(query, cow_ids + [farm_name])
        
    found_cows = cursor.fetchall()
    valid_ids = [c['id'] for c in found_cows]
    
    if valid_ids:
        del_placeholders = ','.join('?' * len(valid_ids))
        cursor.execute(f"DELETE FROM cows WHERE id IN ({del_placeholders})", valid_ids)
        conn.commit()
        
        log_audit_event(
            username=session.get('username'),
            farm_name=farm_name or 'Admin',
            action='BULK_DELETE',
            entity_type='COW',
            details=f"Массовое удаление {len(valid_ids)} записей",
            user_id=session.get('user_id'),
            ip_address=request.remote_addr
        )
        
    conn.close()
    return jsonify({'success': True, 'deleted_count': len(valid_ids)})

@app.route('/add_cows_dynamic')
@login_required
def add_cows_dynamic():
    if session.get('user_type') == 'admin':
        flash('Администраторы не могут добавлять коров', 'warning')
        return redirect(url_for('dashboard'))
        
    return render_template('add_cows_dynamic.html', 
                           categories=Config.DISPOSAL_CATEGORIES, 
                           farm_name=session.get('farm_name'),
                           breeds=Config.BREEDS,
                           age_groups=Config.AGE_GROUPS,
                           today=datetime.now().strftime('%Y-%m-%d'))

@app.route('/add_multiple_cows', methods=['GET', 'POST'])
@login_required
def add_multiple_cows():
    if session.get('user_type') == 'admin':
        if request.is_json:
            return jsonify({'error': 'Администраторы не могут добавлять коров'}), 403
        flash('Администраторы не могут добавлять коров', 'warning')
        return redirect(url_for('dashboard'))
        
    farm_name = session.get('farm_name')
    
    if request.method == 'POST':
        cows_data = []
        if request.is_json:
            data = request.get_json()
            if isinstance(data, list):
                cows_data = data
            elif isinstance(data, dict):
                cows_data = data.get('cows_data', [])
        else:
            raw = request.form.get('cows_data')
            if raw:
                try:
                    parsed = json.loads(raw)
                    cows_data = parsed if isinstance(parsed, list) else parsed.get('cows_data', [])
                except:
                    pass
                    
        conn = get_db_connection()
        cursor = conn.cursor()
        today = datetime.now().strftime('%Y-%m-%d')
        added_count = 0
        
        for item in cows_data:
            cat = item.get('category')
            rea = item.get('reason')
            d_date = item.get('disposal_date')
            ear_tag = item.get('ear_tag') or None
            lactation = item.get('lactation') or None
            weight = item.get('weight') or None
            notes = item.get('notes') or None
            age_group = item.get('age_group') or 'Коровы дойного стада'
            breed = item.get('breed') or 'Черно-пёстрая'
            milk_yield = item.get('milk_yield') or None
            book_value = item.get('book_value') or None
            
            if cat and rea and d_date and d_date <= today:
                cow_id = generate_cow_id(farm_name)
                cursor.execute('''
                    INSERT INTO cows (cow_id, ear_tag, farm_name, category, reason, disposal_date, 
                                      lactation, weight, notes, age_group, breed, milk_yield, book_value, created_by)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                ''', (cow_id, ear_tag, farm_name, cat, rea, d_date, 
                      lactation, weight, notes, age_group, breed, milk_yield, book_value, session.get('user_id')))
                added_count += 1
                
        conn.commit()
        conn.close()
        
        if added_count > 0:
            log_audit_event(
                username=session.get('username'),
                farm_name=farm_name,
                action='CREATE_BATCH',
                entity_type='COW',
                details=f"Пакетно добавлено {added_count} коров",
                user_id=session.get('user_id'),
                ip_address=request.remote_addr
            )
        
        if request.is_json:
            return jsonify({'success': True, 'added_count': added_count})
            
        flash(f'Успешно добавлено {added_count} коров', 'success')
        return redirect(url_for('view_cows'))
        
    return render_template('add_multiple_cows.html', categories=Config.DISPOSAL_CATEGORIES, farm_name=farm_name)

# -------------------- ИМПОРТ ИЗ EXCEL / CSV --------------------

@app.route('/import_cows', methods=['GET', 'POST'])
@login_required
def import_cows():
    if session.get('user_type') == 'admin':
        flash('Только пользователи ферм могут импортировать данные своего хозяйства', 'warning')
        return redirect(url_for('dashboard'))
        
    farm_name = session.get('farm_name')
    today = datetime.now().strftime('%Y-%m-%d')
    
    if request.method == 'POST':
        if 'file' not in request.files:
            flash('Файл не выбран', 'danger')
            return redirect(request.url)
            
        file = request.files['file']
        if file.filename == '':
            flash('Файл не выбран', 'danger')
            return redirect(request.url)
            
        filename = file.filename.lower()
        rows_to_insert = []
        errors = []
        
        try:
            if filename.endswith('.csv'):
                stream = io.StringIO(file.stream.read().decode("utf-8-sig"), newline=None)
                csv_reader = csv.reader(stream, delimiter=';' if ';' in stream.getvalue() else ',')
                header = next(csv_reader, None)
                for line_num, row in enumerate(csv_reader, start=2):
                    if not row or len(row) < 3:
                        continue
                    cat = row[0].strip().lower()
                    rea = row[1].strip()
                    d_date = row[2].strip()
                    ear_tag = row[3].strip() if len(row) > 3 and row[3].strip() else None
                    lact = int(row[4]) if len(row) > 4 and row[4].strip().isdigit() else None
                    w = float(row[5].replace(',', '.')) if len(row) > 5 and row[5].strip() else None
                    notes = row[6].strip() if len(row) > 6 and row[6].strip() else None
                    
                    if cat not in ['падёж', 'выбраковка', 'санитарный']:
                        errors.append(f"Строка {line_num}: недопустимая категория '{cat}'")
                        continue
                    if d_date > today:
                        errors.append(f"Строка {line_num}: дата {d_date} позже сегодняшней")
                        continue
                    rows_to_insert.append((cat, rea, d_date, ear_tag, lact, w, notes))
                    
            elif filename.endswith(('.xlsx', '.xls')):
                wb = openpyxl.load_workbook(file, data_only=True)
                ws = wb.active
                for line_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                    if not row or not row[0]:
                        continue
                    cat = str(row[0]).strip().lower()
                    rea = str(row[1]).strip() if row[1] else ''
                    
                    raw_date = row[2]
                    if isinstance(raw_date, datetime):
                        d_date = raw_date.strftime('%Y-%m-%d')
                    else:
                        d_date = str(raw_date).strip()
                        
                    ear_tag = str(row[3]).strip() if len(row) > 3 and row[3] else None
                    lact = int(row[4]) if len(row) > 4 and str(row[4]).strip().isdigit() else None
                    try:
                        w = float(str(row[5]).replace(',', '.')) if len(row) > 5 and row[5] else None
                    except ValueError:
                        w = None
                    notes = str(row[6]).strip() if len(row) > 6 and row[6] else None
                    
                    if cat not in ['падёж', 'выбраковка', 'санитарный']:
                        errors.append(f"Строка {line_num}: недопустимая категория '{cat}'")
                        continue
                    if d_date > today:
                        errors.append(f"Строка {line_num}: дата {d_date} позже сегодняшней")
                        continue
                    rows_to_insert.append((cat, rea, d_date, ear_tag, lact, w, notes))
            else:
                flash('Поддерживаются только форматы .xlsx, .xls и .csv', 'danger')
                return redirect(request.url)
                
        except Exception as e:
            flash(f'Ошибка обработки файла: {str(e)}', 'danger')
            return redirect(request.url)
            
        if rows_to_insert:
            conn = get_db_connection()
            cursor = conn.cursor()
            inserted_count = 0
            for r in rows_to_insert:
                cow_id = generate_cow_id(farm_name)
                cursor.execute('''
                    INSERT INTO cows (cow_id, ear_tag, farm_name, category, reason, disposal_date, lactation, weight, notes, created_by)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                ''', (cow_id, r[3], farm_name, r[0], r[1], r[2], r[4], r[5], r[6], session.get('user_id')))
                inserted_count += 1
            conn.commit()
            conn.close()
            
            log_audit_event(
                username=session.get('username'),
                farm_name=farm_name,
                action='IMPORT',
                entity_type='COW',
                details=f"Импортировано {inserted_count} записей из файла {file.filename}",
                user_id=session.get('user_id'),
                ip_address=request.remote_addr
            )
            
            flash(f'Успешно импортировано {inserted_count} коров!', 'success')
            if errors:
                flash(f'Внимание: пропущено строк с ошибками: {len(errors)}', 'warning')
            return redirect(url_for('view_cows'))
        else:
            flash('В файле не найдено корректных строк для импорта.', 'danger')
            return redirect(request.url)
            
    return render_template('import_cows.html', farm_name=farm_name)

@app.route('/download_import_template')
@login_required
def download_import_template():
    """Скачать образец Excel-шаблона для импорта"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Шаблон_Импорта"
    
    headers = ['Категория (падёж/выбраковка/санитарный)*', 'Причина выбытия*', 'Дата (ГГГГ-ММ-ДД)*', 'Инв./Ушной №', 'Лактация', 'Масса (кг)', 'Примечание']
    ws.append(headers)
    
    # Примерные строки
    ws.append(['падёж', 'Острая тимпания', datetime.now().strftime('%Y-%m-%d'), '7701', 3, 520.5, 'Заключение ветврача'])
    ws.append(['выбраковка', 'Агалактия', datetime.now().strftime('%Y-%m-%d'), '7702', 4, 480.0, 'Низкая продуктивность'])
    ws.append(['санитарный', 'Травма конечностей', datetime.now().strftime('%Y-%m-%d'), '7703', 2, 510.0, 'Травма в секции'])
    
    # Стилизация шапки
    for cell in ws[1]:
        cell.font = openpyxl.styles.Font(bold=True, color="FFFFFF")
        cell.fill = openpyxl.styles.PatternFill(start_color="3B82F6", end_color="3B82F6", fill_type="solid")
        cell.alignment = openpyxl.styles.Alignment(horizontal="center")
        
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 16)
        
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return send_file(out, download_name="shablon_importa_korov.xlsx", as_attachment=True, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

# -------------------- УПРАВЛЕНИЕ СПРАВОЧНИКОМ ПРИЧИН --------------------

@app.route('/reasons', methods=['GET', 'POST'])
@admin_required
def manage_reasons():
    """Управление справочником причин выбытия скота (для администратора)"""
    conn = get_db_connection()
    cursor = conn.cursor()
    
    if request.method == 'POST':
        category = request.form.get('category')
        name = request.form.get('name', '').strip()
        
        if category and name:
            cursor.execute('INSERT OR IGNORE INTO disposal_reasons (category, name, is_active) VALUES (?, ?, 1)', (category, name))
            conn.commit()
            
            log_audit_event(
                username=session.get('username'),
                farm_name='Admin',
                action='CREATE_REASON',
                entity_type='REASON',
                details=f"Добавлена новая причина '{name}' в категорию '{category}'",
                user_id=session.get('user_id'),
                ip_address=request.remote_addr
            )
            flash(f"Причина '{name}' успешно добавлена в справочник", 'success')
            
    cursor.execute('SELECT * FROM disposal_reasons ORDER BY category, name')
    reasons = [dict(r) for r in cursor.fetchall()]
    conn.close()
    
    return render_template('manage_reasons.html', reasons=reasons)

@app.route('/reasons/toggle/<int:reason_id>', methods=['POST'])
@admin_required
def toggle_reason(reason_id):
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute('SELECT * FROM disposal_reasons WHERE id = ?', (reason_id,))
    reason = cursor.fetchone()
    if reason:
        new_status = 0 if reason['is_active'] == 1 else 1
        cursor.execute('UPDATE disposal_reasons SET is_active = ? WHERE id = ?', (new_status, reason_id))
        conn.commit()
        status_text = 'активирована' if new_status == 1 else 'скрыта'
        flash(f"Причина '{reason['name']}' {status_text}", 'info')
    conn.close()
    return redirect(url_for('manage_reasons'))

@app.route('/reasons/delete/<int:reason_id>', methods=['POST'])
@admin_required
def delete_reason(reason_id):
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute('SELECT * FROM disposal_reasons WHERE id = ?', (reason_id,))
    reason = cursor.fetchone()
    if reason:
        cursor.execute('DELETE FROM disposal_reasons WHERE id = ?', (reason_id,))
        conn.commit()
        log_audit_event(
            username=session.get('username'),
            farm_name='Admin',
            action='DELETE_REASON',
            entity_type='REASON',
            details=f"Удалена причина '{reason['name']}' из категории '{reason['category']}'",
            user_id=session.get('user_id'),
            ip_address=request.remote_addr
        )
        flash(f"Причина '{reason['name']}' удалена из справочника", 'success')
    conn.close()
    return redirect(url_for('manage_reasons'))

# -------------------- ЖУРНАЛ АУДИТА (AUDIT LOG) --------------------

@app.route('/audit_log')
@admin_required
def audit_log():
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute('SELECT * FROM audit_logs ORDER BY created_at DESC LIMIT 200')
    logs = [dict(r) for r in cursor.fetchall()]
    conn.close()
    return render_template('audit_log.html', logs=logs)

# -------------------- УПРАВЛЕНИЕ ПОЛЬЗОВАТЕЛЯМИ --------------------

@app.route('/register', methods=['GET', 'POST'])
@admin_required
def register():
    if request.method == 'POST':
        username = request.form.get('username', '').strip()
        password = request.form.get('password', '')
        user_type = request.form.get('user_type')
        farm_name = request.form.get('farm_name', '').strip() if user_type == 'farm' else None
        
        success, message = register_user(username, password, user_type, farm_name)
        if success:
            log_audit_event(
                username=session.get('username'),
                farm_name='Admin',
                action='CREATE_USER',
                entity_type='USER',
                details=f"Зарегистрирован пользователь {username} ({user_type}, ферма: {farm_name or '-'})",
                user_id=session.get('user_id'),
                ip_address=request.remote_addr
            )
            flash(message, 'success')
            return redirect(url_for('user_management'))
        else:
            flash(message, 'danger')
            
    return render_template('register.html')

@app.route('/user_management')
@admin_required
def user_management():
    users = get_all_users()
    return render_template('user_management.html', users=users)

@app.route('/edit_user/<int:user_id>', methods=['GET', 'POST'])
@admin_required
def edit_user(user_id):
    user = get_user_by_id(user_id)
    if not user:
        flash('Пользователь не найден', 'danger')
        return redirect(url_for('user_management'))
        
    if request.method == 'POST':
        username = request.form.get('username', '').strip()
        password = request.form.get('password', '')
        user_type = request.form.get('user_type')
        farm_name = request.form.get('farm_name', '').strip() if user_type == 'farm' else None
        
        success, message = update_user(user_id, username, password if password else None, user_type, farm_name)
        if success:
            log_audit_event(
                username=session.get('username'),
                farm_name='Admin',
                action='UPDATE_USER',
                entity_type='USER',
                entity_id=user_id,
                details=f"Обновлен профиль пользователя {username}",
                user_id=session.get('user_id'),
                ip_address=request.remote_addr
            )
            flash(message, 'success')
            return redirect(url_for('user_management'))
        else:
            flash(message, 'danger')
            
    return render_template('edit_user.html', user=user)

@app.route('/delete_user/<int:user_id>', methods=['POST'])
@admin_required
def delete_user_route(user_id):
    if user_id == session.get('user_id'):
        flash('Вы не можете удалить свою собственную учетную запись', 'danger')
        return redirect(url_for('user_management'))
        
    user = get_user_by_id(user_id)
    if user:
        conn = get_db_connection()
        cursor = conn.cursor()
        if user.user_type == 'farm' and user.farm_name:
            cursor.execute('DELETE FROM cows WHERE farm_name = ?', (user.farm_name,))
        cursor.execute('DELETE FROM users WHERE id = ?', (user_id,))
        conn.commit()
        conn.close()
        
        log_audit_event(
            username=session.get('username'),
            farm_name='Admin',
            action='DELETE_USER',
            entity_type='USER',
            entity_id=user_id,
            details=f"Удален пользователь {user.username} (ферма: {user.farm_name or '-'})",
            user_id=session.get('user_id'),
            ip_address=request.remote_addr
        )
        flash(f'Пользователь {user.username} успешно удален', 'success')
    else:
        flash('Пользователь не найден', 'danger')
        
    return redirect(url_for('user_management'))

# -------------------- ОТЧЁТЫ И ЭКСПОРТ --------------------

@app.route('/reports')
@login_required
def reports():
    return render_template('reports.html', user_type=session.get('user_type'))

@app.route('/generate_report', methods=['POST'])
@login_required
def generate_report():
    report_type = request.form.get('report_type')
    report_format = request.form.get('format')
    start_date = request.form.get('start_date')
    end_date = request.form.get('end_date')
    farm_filter = request.form.get('farm_filter')
    report_category = request.form.get('report_category')
    
    if session.get('user_type') == 'farm':
        farm_filter = session.get('farm_name')
        
    if report_type in ['form_209_apk', 'sp54']:
        wb = generate_form_209_apk_excel(start_date, end_date, farm_filter)
        filename = f"akt_209_apk_RB_{start_date}_{end_date}.xlsx"
        out = io.BytesIO()
        wb.save(out)
        out.seek(0)
        return send_file(out, download_name=filename, as_attachment=True, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

    if report_type == 'form_210_apk':
        if report_format == 'preview':
            conn = get_db_connection()
            cursor = conn.cursor()
            where = "WHERE disposal_date BETWEEN ? AND ? AND category IN ('выбраковка', 'санитарный')"
            params = [start_date, end_date]
            if farm_filter and farm_filter != 'all':
                where += " AND farm_name = ?"
                params.append(farm_filter)
            cursor.execute(f"SELECT * FROM cows {where} ORDER BY disposal_date ASC", params)
            cows_list = [dict(r) for r in cursor.fetchall()]
            conn.close()
            
            total_weight = sum(float(c['weight'] or 0) for c in cows_list)
            total_value = sum(float(c['book_value'] or 0) for c in cows_list)
            org_name = get_setting('farm_org_name', 'ОАО «Новая Припять»')
            return render_template('form_210_preview.html', cows=cows_list, total_weight=f"{total_weight:.1f}", 
                                   total_value=f"{total_value:.2f}", start_date=start_date, end_date=end_date, 
                                   farm_filter=farm_filter, org_name=org_name)
        else:
            wb = generate_form_210_apk_excel(start_date, end_date, farm_filter)
            filename = f"akt_210_apk_RB_{start_date}_{end_date}.xlsx"
            out = io.BytesIO()
            wb.save(out)
            out.seek(0)
            return send_file(out, download_name=filename, as_attachment=True, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

    if report_type == 'aits_registry':
        if report_format == 'xml':
            xml_data = generate_aits_xml_registry(start_date, end_date, farm_filter)
            return Response(
                xml_data,
                mimetype="application/xml; charset=utf-8",
                headers={"Content-disposition": f"attachment; filename=aits_registry_{start_date}_{end_date}.xml"}
            )
        else:
            csv_data = generate_aits_csv_registry(start_date, end_date, farm_filter)
            return Response(
                '\ufeff' + csv_data,
                mimetype="text/csv; charset=utf-8",
                headers={"Content-disposition": f"attachment; filename=aits_registry_{start_date}_{end_date}.csv"}
            )

    if report_format == 'csv':
        csv_data = generate_csv_report(start_date, end_date, farm_filter)
        return Response(
            '\ufeff' + csv_data,
            mimetype="text/csv; charset=utf-8",
            headers={"Content-disposition": f"attachment; filename=report_{start_date}_{end_date}.csv"}
        )
        
    elif report_format == 'preview':
        conn = get_db_connection()
        cursor = conn.cursor()
        where = "WHERE disposal_date BETWEEN ? AND ?"
        params = [start_date, end_date]
        if farm_filter and farm_filter != 'all':
            where += " AND farm_name = ?"
            params.append(farm_filter)
        cursor.execute(f"SELECT * FROM cows {where} ORDER BY disposal_date DESC", params)
        cows_list = [dict(r) for r in cursor.fetchall()]
        conn.close()
        
        summary = {
            'total_cows': len(cows_list),
            'by_category': {
                'падёж': sum(1 for c in cows_list if c['category'] == 'падёж'),
                'выбраковка': sum(1 for c in cows_list if c['category'] == 'выбраковка'),
                'санитарный': sum(1 for c in cows_list if c['category'] == 'санитарный')
            },
            'total_weight': sum(float(c['weight'] or 0) for c in cows_list)
        }
        return render_template('report_preview.html', data=cows_list, summary=summary, start_date=start_date, end_date=end_date, farm_filter=farm_filter)
        
    else:  # Excel
        if report_type == 'matrix':
            wb = generate_matrix_excel_report(report_category or 'падёж', start_date, end_date)
            filename = f"matrix_{report_category}_{start_date}_{end_date}.xlsx"
        elif report_type == 'structured':
            wb = generate_structured_excel_report(start_date, end_date, farm_filter)
            filename = f"structured_report_{start_date}_{end_date}.xlsx"
        else:
            wb = generate_simple_excel_report(start_date, end_date, farm_filter)
            filename = f"simple_report_{start_date}_{end_date}.xlsx"
            
        out = io.BytesIO()
        wb.save(out)
        out.seek(0)
        return send_file(out, download_name=filename, as_attachment=True, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

@app.route('/generate_original_excel', methods=['POST'])
@admin_required
def generate_original_excel_route():
    wb = generate_official_template_excel()
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return send_file(out, download_name=f"svodka_{datetime.now().strftime('%Y-%m-%d')}.xlsx", as_attachment=True, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

# -------------------- ЭКОНОМИЧЕСКИЙ АНАЛИЗ (BYN) --------------------

@app.route('/economics')
@login_required
def economics_dashboard():
    user_type = session.get('user_type')
    farm_name = request.args.get('farm_name', 'all')
    if user_type == 'farm':
        farm_name = session.get('farm_name')
        
    today = datetime.now()
    first_day = datetime(today.year, today.month, 1).strftime('%Y-%m-%d')
    today_str = today.strftime('%Y-%m-%d')
    
    start_date = request.args.get('start_date', first_day)
    end_date = request.args.get('end_date', today_str)
    
    econ_data = calculate_economic_losses(start_date, end_date, farm_name)
    
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute('SELECT DISTINCT farm_name FROM users WHERE farm_name IS NOT NULL AND farm_name != "" ORDER BY farm_name')
    farms = [r['farm_name'] for r in cursor.fetchall()]
    conn.close()
    
    unp = get_setting('farm_unp', '190000000')
    org_name = get_setting('farm_org_name', 'ОАО «Новая Припять»')
    
    return render_template('economics.html', data=econ_data, start_date=start_date, end_date=end_date, 
                           farm_name=farm_name, farms=farms, unp=unp, org_name=org_name)

@app.route('/save_economic_settings', methods=['POST'])
@admin_required
def save_economic_settings():
    meat_price = request.form.get('meat_price_per_kg', '6.20')
    milk_price = request.form.get('milk_price_per_kg', '1.18')
    replacement_cost = request.form.get('replacement_cost', '2950.00')
    farm_unp = request.form.get('farm_unp', '190000000').strip()
    farm_org_name = request.form.get('farm_org_name', 'Сельхозорганизация РБ').strip()
    
    set_setting('meat_price_per_kg', meat_price, 'Закупочная цена 1 кг живой массы КРС (BYN)')
    set_setting('milk_price_per_kg', milk_price, 'Закупочная цена 1 кг базисного молока (BYN)')
    set_setting('replacement_cost', replacement_cost, 'Стоимость восстановления головы / нетели (BYN)')
    set_setting('farm_unp', farm_unp, 'УНП сельхозорганизации (для ГИС AITS)')
    set_setting('farm_org_name', farm_org_name, 'Полное наименование хозяйства (РБ)')
    
    log_audit_event(
        username=session.get('username'),
        farm_name='Admin',
        action='UPDATE',
        entity_type='SETTINGS',
        details=f"Обновлены нормативы цен: мясо {meat_price} руб, молоко {milk_price} руб, нетель {replacement_cost} руб",
        user_id=session.get('user_id'),
        ip_address=request.remote_addr
    )
    flash('Экономические нормативы и реквизиты хозяйства успешно сохранены', 'success')
    return redirect(url_for('economics_dashboard'))

# -------------------- ПАТОЛОГОАНАТОМИЧЕСКОЕ ВСКРЫТИЕ (РБ) --------------------

@app.route('/autopsy/<int:cow_id>', methods=['GET'])
@login_required
def autopsy_route(cow_id):
    cow = get_cow_for_autopsy(cow_id)
    if not cow:
        flash('Животное не найдено', 'danger')
        return redirect(url_for('view_cows'))
        
    protocol = parse_autopsy_protocol(cow.get('autopsy_protocol'))
    return render_template('autopsy_form.html', cow=cow, protocol=protocol)

@app.route('/save_autopsy/<int:cow_id>', methods=['POST'])
@login_required
def save_autopsy_route(cow_id):
    cow = get_cow_for_autopsy(cow_id)
    if not cow:
        flash('Животное не найдено', 'danger')
        return redirect(url_for('view_cows'))
        
    data = {
        'farm_name': cow.get('farm_name'),
        'autopsy_vet': request.form.get('autopsy_vet'),
        'autopsy_date': request.form.get('autopsy_date'),
        'autopsy_lab_sample': request.form.get('autopsy_lab_sample'),
        'anamnesis': request.form.get('anamnesis'),
        'external_exam': request.form.get('external_exam'),
        'respiratory': request.form.get('respiratory'),
        'cardiovascular': request.form.get('cardiovascular'),
        'digestive': request.form.get('digestive'),
        'liver_spleen': request.form.get('liver_spleen'),
        'pat_diagnosis': request.form.get('pat_diagnosis'),
        'conclusion': request.form.get('conclusion'),
        'lab_tests': request.form.get('lab_tests'),
        'lab_doc_num': request.form.get('lab_doc_num')
    }
    
    save_autopsy_protocol(cow_id, data, session.get('username'), request.remote_addr)
    flash(f"Протокол вскрытия для животного {cow['cow_id']} успешно сохранен", 'success')
    return redirect(url_for('autopsy_route', cow_id=cow_id))

@app.route('/autopsy/print/<int:cow_id>')
@login_required
def print_autopsy(cow_id):
    cow = get_cow_for_autopsy(cow_id)
    if not cow:
        flash('Животное не найдено', 'danger')
        return redirect(url_for('view_cows'))
        
    protocol = parse_autopsy_protocol(cow.get('autopsy_protocol'))
    org_name = get_setting('farm_org_name', 'ОАО «Новая Припять»')
    return render_template('autopsy_print.html', cow=cow, protocol=protocol, org_name=org_name)

# -------------------- РЕЗЕРВНОЕ КОПИРОВАНИЕ И ВОССТАНОВЛЕНИЕ (BACKUP) --------------------

@app.route('/backup')
@admin_required
def backup_page():
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM audit_logs WHERE action = 'BACKUP' ORDER BY created_at DESC LIMIT 20")
    backup_logs = [dict(r) for r in cursor.fetchall()]
    conn.close()
    return render_template('backup.html', backup_logs=backup_logs)

@app.route('/backup/download')
@admin_required
def download_db_backup():
    file_bytes, filename = create_db_backup_bytes()
    log_audit_event(
        username=session.get('username'),
        farm_name='Admin',
        action='BACKUP',
        entity_type='SETTINGS',
        details=f"Скачана резервная копия базы данных: {filename}",
        user_id=session.get('user_id'),
        ip_address=request.remote_addr
    )
    return send_file(
        io.BytesIO(file_bytes),
        download_name=filename,
        as_attachment=True,
        mimetype='application/octet-stream'
    )

@app.route('/backup/restore', methods=['POST'])
@admin_required
def restore_db_backup():
    if 'backup_file' not in request.files:
        flash('Файл резервной копии не выбран', 'danger')
        return redirect(url_for('backup_page'))
        
    file = request.files['backup_file']
    if file.filename == '':
        flash('Файл резервной копии не выбран', 'danger')
        return redirect(url_for('backup_page'))
        
    success, msg = restore_db_from_upload(file, session.get('username'), request.remote_addr)
    if success:
        flash(msg, 'success')
    else:
        flash(msg, 'danger')
    return redirect(url_for('backup_page'))

@app.route('/seed_test_data')
@login_required
def seed_test_data():
    from add_test_cows import generate_test_cows
    generate_test_cows(1000)
    flash('Успешно добавлено 1 000 коров в базу данных!', 'success')
    return redirect(url_for('view_cows'))

# -------------------- API МАРШРУТЫ --------------------

@app.route('/get_farms')
@login_required
def get_farms():
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute('SELECT DISTINCT farm_name FROM users WHERE farm_name IS NOT NULL AND farm_name != "" ORDER BY farm_name')
    farms = [r['farm_name'] for r in cursor.fetchall()]
    conn.close()
    return jsonify(farms)

@app.route('/get_reasons/<path:category>')
@login_required
def get_reasons(category):
    reasons = get_reasons_for_category(category)
    return jsonify(reasons)

if __name__ == '__main__':
    init_db()
    from create_users import create_default_users
    create_default_users()
    app.run(host='0.0.0.0', port=3000, debug=True)
