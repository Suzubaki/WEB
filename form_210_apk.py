# form_210_apk.py - Генерация Типовой формы 210-АПК (Республика Беларусь)
import sqlite3
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime
from database import get_db_connection, get_setting

def generate_form_210_apk_excel(start_date, end_date, farm_name=None):
    """
    Генерация Акта на выбраковку животных из основного стада
    (Типовая форма 210-АПК, утвержденная Министерством сельского хозяйства и продовольствия Республики Беларусь)
    """
    conn = get_db_connection()
    cursor = conn.cursor()
    
    # Для формы 210-АПК отбираем только категорию 'выбраковка' (и санитарную прирезку основного стада)
    where_clause = "WHERE disposal_date BETWEEN ? AND ? AND category IN ('выбраковка', 'санитарный')"
    params = [start_date, end_date]
    
    if farm_name and farm_name != 'all':
        where_clause += " AND farm_name = ?"
        params.append(farm_name)
        
    cursor.execute(f'''
        SELECT cow_id, ear_tag, farm_name, category, reason, disposal_date, 
               lactation, weight, notes, age_group, breed, milk_yield, book_value 
        FROM cows 
        {where_clause}
        ORDER BY disposal_date ASC
    ''', params)
    
    rows = cursor.fetchall()
    conn.close()
    
    org_name = get_setting('farm_org_name', 'Сельскохозяйственная организация (РБ)')
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Форма 210-АПК (РБ)"
    
    # Стили
    title_font = Font(name='Arial', size=11, bold=True)
    header_font = Font(name='Arial', size=9, bold=True)
    body_font = Font(name='Arial', size=9)
    bold_body_font = Font(name='Arial', size=9, bold=True)
    small_font = Font(name='Arial', size=8, italic=True)
    
    thin_border = Border(
        left=Side(style='thin', color='000000'),
        right=Side(style='thin', color='000000'),
        top=Side(style='thin', color='000000'),
        bottom=Side(style='thin', color='000000')
    )
    header_fill = PatternFill(start_color="DCE7F5", end_color="DCE7F5", fill_type="solid")
    
    # Правый верхний угол: гриф формы
    ws.merge_cells('H1:K1')
    ws['H1'] = "Типовая форма 210-АПК"
    ws['H1'].font = Font(name='Arial', size=9, bold=True)
    ws['H1'].alignment = Alignment(horizontal='right', vertical='center')
    
    ws.merge_cells('G2:K2')
    ws['G2'] = "Утверждена постановлением Минсельхозпрода Республики Беларусь"
    ws['G2'].font = small_font
    ws['G2'].alignment = Alignment(horizontal='right', vertical='center')
    
    # Левый верхний угол: реквизиты
    ws['A4'] = f"Сельхозорганизация: {org_name}"
    ws['A4'].font = Font(name='Arial', size=10, bold=True)
    ws['A5'] = f"Ферма / МТФ: {farm_name if (farm_name and farm_name != 'all') else 'Сводный акт по хозяйству'}"
    ws['A5'].font = Font(name='Arial', size=9, bold=True)
    
    # Гриф УТВЕРЖДАЮ
    ws.merge_cells('H4:K4')
    ws['H4'] = "УТВЕРЖДАЮ"
    ws['H4'].font = Font(name='Arial', size=10, bold=True)
    ws['H4'].alignment = Alignment(horizontal='center', vertical='center')
    
    ws.merge_cells('H5:K5')
    ws['H5'] = "Руководитель организации (Директор): ________________"
    ws['H5'].font = body_font
    ws['H5'].alignment = Alignment(horizontal='center', vertical='center')
    
    ws.merge_cells('A7:K7')
    ws['A7'] = "АКТ НА ВЫБРАКОВКУ ЖИВОТНЫХ ИЗ ОСНОВНОГО СТАДА"
    ws['A7'].font = title_font
    ws['A7'].alignment = Alignment(horizontal='center', vertical='center')
    
    ws.merge_cells('A8:K8')
    ws['A8'] = f"за период с {start_date} по {end_date} года"
    ws['A8'].font = body_font
    ws['A8'].alignment = Alignment(horizontal='center', vertical='center')
    
    # Табличная часть (11 колонок)
    headers = [
        ('№ п/п', 5),
        ('Идент. № AITS / Инв. №', 18),
        ('Подразделение (МТФ)', 16),
        ('Порода', 18),
        ('Лактация / Возраст', 14),
        ('Живая масса (кг)', 13),
        ('Удой за посл. лактацию (кг)', 16),
        ('Балансовая стоим. (BYN)', 15),
        ('Причина выбраковки (диагноз)', 24),
        ('Вид выбытия', 13),
        ('Направление использования', 20)
    ]
    
    row_num = 10
    for col_idx, (h_title, width) in enumerate(headers, 1):
        cell = ws.cell(row=row_num, column=col_idx, value=h_title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = thin_border
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = width
    ws.row_dimensions[row_num].height = 32
    
    total_weight = 0.0
    total_value = 0.0
    
    for idx, r in enumerate(rows, 1):
        row_num += 1
        ws.row_dimensions[row_num].height = 22
        
        w_val = float(r['weight'] or 0.0)
        total_weight += w_val
        
        b_val = float(r['book_value'] or 0.0)
        total_value += b_val
        
        usage = "Сдача на мясокомбинат" if r['category'] == 'выбраковка' else "Санитарная бойня"
        
        vals = [
            idx,
            r['ear_tag'] or r['cow_id'],
            r['farm_name'],
            r['breed'] or 'Черно-пёстрая',
            r['lactation'] or '-',
            f"{w_val:.1f}" if w_val > 0 else "-",
            f"{r['milk_yield']:.0f}" if r['milk_yield'] else "-",
            f"{b_val:.2f}" if b_val > 0 else "-",
            r['reason'],
            r['category'].capitalize(),
            usage
        ]
        
        for col_idx, val in enumerate(vals, 1):
            cell = ws.cell(row=row_num, column=col_idx, value=val)
            cell.font = body_font
            cell.border = thin_border
            if col_idx in [1, 2, 5, 6, 7, 8, 10]:
                cell.alignment = Alignment(horizontal='center', vertical='center')
            else:
                cell.alignment = Alignment(horizontal='left', vertical='center')
                
    # Строка Итого
    row_num += 1
    ws.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=5)
    ws.cell(row=row_num, column=1, value="ИТОГО ВЫБРАКОВАНО:").font = bold_body_font
    ws.cell(row=row_num, column=1).alignment = Alignment(horizontal='right', vertical='center')
    for c in range(1, 6):
        ws.cell(row=row_num, column=c).border = thin_border
        ws.cell(row=row_num, column=c).fill = header_fill
        
    c_weight = ws.cell(row=row_num, column=6, value=f"{total_weight:.1f} кг" if total_weight > 0 else "-")
    c_weight.font = bold_body_font
    c_weight.alignment = Alignment(horizontal='center', vertical='center')
    c_weight.border = thin_border
    c_weight.fill = header_fill
    
    c_empty7 = ws.cell(row=row_num, column=7, value="")
    c_empty7.border = thin_border
    c_empty7.fill = header_fill
    
    c_value = ws.cell(row=row_num, column=8, value=f"{total_value:.2f} руб." if total_value > 0 else "-")
    c_value.font = bold_body_font
    c_value.alignment = Alignment(horizontal='center', vertical='center')
    c_value.border = thin_border
    c_value.fill = header_fill
    
    ws.merge_cells(start_row=row_num, start_column=9, end_row=row_num, end_column=11)
    c_count = ws.cell(row=row_num, column=9, value=f"{len(rows)} голов")
    c_count.font = bold_body_font
    c_count.alignment = Alignment(horizontal='center', vertical='center')
    for c in range(9, 12):
        ws.cell(row=row_num, column=c).border = thin_border
        ws.cell(row=row_num, column=c).fill = header_fill
        
    # Блок подписей комиссии сельхозорганизации
    row_num += 3
    ws.cell(row=row_num, column=1, value="Комиссия по выбраковке животных:").font = bold_body_font
    row_num += 1
    ws.cell(row=row_num, column=1, value="Главный зоотехник:          _______________________ (подпись)").font = body_font
    row_num += 1
    ws.cell(row=row_num, column=1, value="Главный ветеринарный врач: _______________________ (подпись)").font = body_font
    row_num += 1
    ws.cell(row=row_num, column=1, value="Главный бухгалтер:          _______________________ (подпись)").font = body_font
    row_num += 1
    ws.cell(row=row_num, column=1, value="Заведующий МТФ / фермой:    _______________________ (подпись)").font = body_font
    row_num += 1
    ws.cell(row=row_num, column=1, value="Материально ответственное:  _______________________ (подпись)").font = body_font
    
    return wb
