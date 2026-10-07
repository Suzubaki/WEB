# routes/cows.py - Управление записями выбытия скота (Журнал коров)
import os
import io
import csv
import json
from datetime import datetime
from flask import (
    Blueprint, render_template, request, redirect, 
    url_for, flash, session, jsonify, send_file
)
import openpyxl

from config import Config
from database import get_db_connection, generate_cow_id, log_audit_event
from .auth import login_required

cows_bp = Blueprint('cows', __name__)

@cows_bp.route('/cows')
@login_required
def view_cows():
    """Просмотр журнала выбытия коров"""
    user_type = session.get('user_type')
    farm_name = session.get('farm_name')
    
    conn = get_db_connection()
    cursor = conn.cursor()
    
    if user_type == 'admin':
        cursor.execute('SELECT * FROM cows ORDER BY disposal_date DESC, id DESC')
    else:
        cursor.execute('SELECT * FROM cows WHERE farm_name = ? ORDER BY disposal_date DESC, id DESC', (farm_name,))
        
    cows_data = cursor.fetchall()
    conn.close()
    
    cows = [dict(c) for c in cows_data]
    return render_template('view_cows.html', cows=cows, user_type=user_type)

@cows_bp.route('/add_cow', methods=['GET', 'POST'])
@login_required
def add_cow():
    """Добавление единичной записи выбытия КРС"""
    if session.get('user_type') == 'admin':
        flash('Администраторы не могут добавлять коров', 'warning')
        return redirect(url_for('main.dashboard'))
        
    farm_name = session.get('farm_name')
    
    if request.method == 'POST':
        category = request.form.get('category')
        reason = request.form.get('reason')
        disposal_date = request.form.get('disposal_date')
        ear_tag = request.form.get('ear_tag', '').strip() or None
        lactation = request.form.get('lactation') or None
        weight = request.form.get('weight') or None
        notes = request.form.get('notes', '').strip() or None
        age_group = request.form.get('age_group') or 'Коровы дойного стада'
        breed = request.form.get('breed') or 'Черно-пёстрая'
        milk_yield = request.form.get('milk_yield') or None
        book_value = request.form.get('book_value') or None
        
        if not category or not reason or not disposal_date:
            flash('Заполните обязательные поля: категория, причина и дата', 'danger')
            return render_template('add_cow.html', categories=Config.DISPOSAL_CATEGORIES, farm_name=farm_name, breeds=Config.BREEDS, age_groups=Config.AGE_GROUPS)
            
        today = datetime.now().strftime('%Y-%m-%d')
        if disposal_date > today:
            flash('Дата выбытия не может быть позже сегодняшнего дня', 'danger')
            return render_template('add_cow.html', categories=Config.DISPOSAL_CATEGORIES, farm_name=farm_name, breeds=Config.BREEDS, age_groups=Config.AGE_GROUPS)
            
        cow_id = generate_cow_id(farm_name)
        
        conn = get_db_connection()
        cursor = conn.cursor()
        cursor.execute('''
            INSERT INTO cows (cow_id, ear_tag, farm_name, category, reason, disposal_date, 
                              lactation, weight, notes, age_group, breed, milk_yield, book_value, created_by)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        ''', (cow_id, ear_tag, farm_name, category, reason, disposal_date, 
              lactation, weight, notes, age_group, breed, milk_yield, book_value, session.get('user_id')))
        new_id = cursor.lastrowid
        conn.commit()
        conn.close()
        
        log_audit_event(
            username=session.get('username'),
            farm_name=farm_name,
            action='CREATE',
            entity_type='COW',
            entity_id=new_id,
            details=f"Добавлена корова {cow_id} (Ушной № {ear_tag or '-'}), {category}: {reason}",
            user_id=session.get('user_id'),
            ip_address=request.remote_addr
        )
        
        flash(f'Корова успешно добавлена с системным номером: {cow_id}', 'success')
        return redirect(url_for('cows.view_cows'))
        
    return render_template('add_cow.html', categories=Config.DISPOSAL_CATEGORIES, farm_name=farm_name, breeds=Config.BREEDS, age_groups=Config.AGE_GROUPS)

@cows_bp.route('/edit_cow/<int:cow_id>', methods=['GET', 'POST'])
@login_required
def edit_cow(cow_id):
    """Редактирование записи о корове"""
    user_type = session.get('user_type')
    farm_name = session.get('farm_name')
    
    conn = get_db_connection()
    cursor = conn.cursor()
    
    if user_type == 'admin':
        cursor.execute('SELECT * FROM cows WHERE id = ?', (cow_id,))
    else:
        cursor.execute('SELECT * FROM cows WHERE id = ? AND farm_name = ?', (cow_id, farm_name))
        
    cow_row = cursor.fetchone()
    
    if not cow_row:
        conn.close()
        flash('Запись не найдена или у вас нет прав на её редактирование', 'danger')
        return redirect(url_for('cows.view_cows'))
        
    cow = dict(cow_row)
    
    if request.method == 'POST':
        category = request.form.get('category')
        reason = request.form.get('reason')
        disposal_date = request.form.get('disposal_date')
        ear_tag = request.form.get('ear_tag', '').strip() or None
        lactation = request.form.get('lactation') or None
        weight = request.form.get('weight') or None
        notes = request.form.get('notes', '').strip() or None
        age_group = request.form.get('age_group') or 'Коровы дойного стада'
        breed = request.form.get('breed') or 'Черно-пёстрая'
        milk_yield = request.form.get('milk_yield') or None
        book_value = request.form.get('book_value') or None
        
        today = datetime.now().strftime('%Y-%m-%d')
        if disposal_date > today:
            flash('Дата выбытия не может быть в будущем', 'danger')
            conn.close()
            return render_template('edit_cow.html', cow=cow, categories=Config.DISPOSAL_CATEGORIES)
            
        cursor.execute('''
            UPDATE cows SET 
                category = ?, 
                reason = ?, 
                disposal_date = ?, 
                ear_tag = ?, 
                lactation = ?, 
                weight = ?, 
                notes = ?, 
                age_group = ?,
                breed = ?,
                milk_yield = ?,
                book_value = ?,
                updated_at = CURRENT_TIMESTAMP, 
                updated_by = ?
            WHERE id = ?
        ''', (category, reason, disposal_date, ear_tag, lactation, weight, notes, 
              age_group, breed, milk_yield, book_value, session.get('user_id'), cow_id))
        conn.commit()
        conn.close()
        
        log_audit_event(
            username=session.get('username'),
            farm_name=cow['farm_name'],
            action='UPDATE',
            entity_type='COW',
            entity_id=cow_id,
            details=f"Отредактирована запись {cow['cow_id']}. Новые данные: {category} - {reason}, дата {disposal_date}",
            user_id=session.get('user_id'),
            ip_address=request.remote_addr
        )
        
        flash(f'Запись {cow["cow_id"]} успешно обновлена', 'success')
        return redirect(url_for('cows.view_cows'))
        
    conn.close()
    return render_template('edit_cow.html', cow=cow, categories=Config.DISPOSAL_CATEGORIES)

@cows_bp.route('/delete_cow/<int:cow_id>', methods=['POST'])
@login_required
def delete_cow(cow_id):
    """Удаление записи выбытия"""
    user_type = session.get('user_type')
    farm_name = session.get('farm_name')
    
    conn = get_db_connection()
    cursor = conn.cursor()
    
    if user_type == 'admin':
        cursor.execute('SELECT * FROM cows WHERE id = ?', (cow_id,))
    else:
        cursor.execute('SELECT * FROM cows WHERE id = ? AND farm_name = ?', (cow_id, farm_name))
        
    cow = cursor.fetchone()
    if not cow:
        conn.close()
        if request.is_json:
            return jsonify({'error': 'Запись не найдена'}), 404
        flash('Запись не найдена', 'danger')
        return redirect(url_for('cows.view_cows'))
        
    cow_dict = dict(cow)
    cursor.execute('DELETE FROM cows WHERE id = ?', (cow_id,))
    conn.commit()
    conn.close()
    
    log_audit_event(
        username=session.get('username'),
        farm_name=cow_dict['farm_name'],
        action='DELETE',
        entity_type='COW',
        entity_id=cow_id,
        details=f"Удалена запись {cow_dict['cow_id']} ({cow_dict['category']}: {cow_dict['reason']})",
        user_id=session.get('user_id'),
        ip_address=request.remote_addr
    )
    
    if request.is_json or request.headers.get('X-Requested-With') == 'XMLHttpRequest':
        return jsonify({'success': True, 'message': 'Запись удалена'})
        
    flash(f'Запись {cow_dict["cow_id"]} успешно удалена', 'success')
    return redirect(url_for('cows.view_cows'))

@cows_bp.route('/bulk_delete_cows', methods=['POST'])
@cows_bp.route('/delete_multiple_cows', methods=['POST'], endpoint='delete_multiple_cows')
@login_required
def bulk_delete_cows():
    """Массовое удаление записей"""
    data = request.get_json() or {}
    cow_ids = data.get('ids', []) or data.get('cow_ids', [])
    if not cow_ids:
        return jsonify({'error': 'Не выбрано ни одной записи'}), 400
        
    user_type = session.get('user_type')
    farm_name = session.get('farm_name')
    
    conn = get_db_connection()
    cursor = conn.cursor()
    
    placeholders = ','.join('?' * len(cow_ids))
    if user_type == 'admin':
        query = f"SELECT id, cow_id, farm_name FROM cows WHERE id IN ({placeholders})"
        cursor.execute(query, cow_ids)
    else:
        query = f"SELECT id, cow_id, farm_name FROM cows WHERE id IN ({placeholders}) AND farm_name = ?"
        cursor.execute(query, cow_ids + [farm_name])
        
    found_cows = cursor.fetchall()
    valid_ids = [c['id'] for c in found_cows]
    
    if valid_ids:
        del_placeholders = ','.join('?' * len(valid_ids))
        cursor.execute(f"DELETE FROM cows WHERE id IN ({del_placeholders})", valid_ids)
        conn.commit()
        
        log_audit_event(
            username=session.get('username'),
            farm_name=farm_name or 'Admin',
            action='BULK_DELETE',
            entity_type='COW',
            details=f"Массовое удаление {len(valid_ids)} записей",
            user_id=session.get('user_id'),
            ip_address=request.remote_addr
        )
        
    conn.close()
    return jsonify({'success': True, 'deleted_count': len(valid_ids)})

@cows_bp.route('/add_cows_dynamic')
@login_required
def add_cows_dynamic():
    """Динамическая форма множественного ввода записей"""
    if session.get('user_type') == 'admin':
        flash('Администраторы не могут добавлять коров', 'warning')
        return redirect(url_for('main.dashboard'))
        
    return render_template('add_cows_dynamic.html', 
                           categories=Config.DISPOSAL_CATEGORIES, 
                           farm_name=session.get('farm_name'),
                           breeds=Config.BREEDS,
                           age_groups=Config.AGE_GROUPS,
                           today=datetime.now().strftime('%Y-%m-%d'))

@cows_bp.route('/add_multiple_cows', methods=['GET', 'POST'])
@login_required
def add_multiple_cows():
    """Множественное добавление коров табличным списком"""
    if session.get('user_type') == 'admin':
        if request.is_json:
            return jsonify({'error': 'Администраторы не могут добавлять коров'}), 403
        flash('Администраторы не могут добавлять коров', 'warning')
        return redirect(url_for('main.dashboard'))
        
    farm_name = session.get('farm_name')
    
    if request.method == 'POST':
        cows_data = []
        if request.is_json:
            data = request.get_json()
            if isinstance(data, list):
                cows_data = data
            elif isinstance(data, dict):
                cows_data = data.get('cows_data', [])
        else:
            raw = request.form.get('cows_data')
            if raw:
                try:
                    parsed = json.loads(raw)
                    cows_data = parsed if isinstance(parsed, list) else parsed.get('cows_data', [])
                except:
                    pass
                    
        conn = get_db_connection()
        cursor = conn.cursor()
        today = datetime.now().strftime('%Y-%m-%d')
        added_count = 0
        
        for item in cows_data:
            cat = item.get('category')
            rea = item.get('reason')
            d_date = item.get('disposal_date')
            ear_tag = item.get('ear_tag') or None
            lactation = item.get('lactation') or None
            weight = item.get('weight') or None
            notes = item.get('notes') or None
            age_group = item.get('age_group') or 'Коровы дойного стада'
            breed = item.get('breed') or 'Черно-пёстрая'
            milk_yield = item.get('milk_yield') or None
            book_value = item.get('book_value') or None
            
            if cat and rea and d_date and d_date <= today:
                cow_id = generate_cow_id(farm_name)
                cursor.execute('''
                    INSERT INTO cows (cow_id, ear_tag, farm_name, category, reason, disposal_date, 
                                      lactation, weight, notes, age_group, breed, milk_yield, book_value, created_by)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                ''', (cow_id, ear_tag, farm_name, cat, rea, d_date, 
                      lactation, weight, notes, age_group, breed, milk_yield, book_value, session.get('user_id')))
                added_count += 1
                
        conn.commit()
        conn.close()
        
        if added_count > 0:
            log_audit_event(
                username=session.get('username'),
                farm_name=farm_name,
                action='CREATE_BATCH',
                entity_type='COW',
                details=f"Пакетно добавлено {added_count} коров",
                user_id=session.get('user_id'),
                ip_address=request.remote_addr
            )
        
        if request.is_json:
            return jsonify({'success': True, 'added_count': added_count})
            
        flash(f'Успешно добавлено {added_count} коров', 'success')
        return redirect(url_for('cows.view_cows'))
        
    return render_template('add_multiple_cows.html', categories=Config.DISPOSAL_CATEGORIES, farm_name=farm_name)

@cows_bp.route('/import_cows', methods=['GET', 'POST'])
@login_required
def import_cows():
    """Импорт данных выбытия из Excel / CSV файлов"""
    if session.get('user_type') == 'admin':
        flash('Только пользователи ферм могут импортировать данные своего хозяйства', 'warning')
        return redirect(url_for('main.dashboard'))
        
    farm_name = session.get('farm_name')
    today = datetime.now().strftime('%Y-%m-%d')
    
    if request.method == 'POST':
        if 'file' not in request.files:
            flash('Файл не выбран', 'danger')
            return redirect(request.url)
            
        file = request.files['file']
        if file.filename == '':
            flash('Файл не выбран', 'danger')
            return redirect(request.url)
            
        filename = file.filename.lower()
        rows_to_insert = []
        errors = []
        
        try:
            if filename.endswith('.csv'):
                stream = io.StringIO(file.stream.read().decode("utf-8-sig"), newline=None)
                csv_reader = csv.reader(stream, delimiter=';' if ';' in stream.getvalue() else ',')
                header = next(csv_reader, None)
                for line_num, row in enumerate(csv_reader, start=2):
                    if not row or len(row) < 3:
                        continue
                    cat = row[0].strip().lower()
                    rea = row[1].strip()
                    d_date = row[2].strip()
                    ear_tag = row[3].strip() if len(row) > 3 and row[3].strip() else None
                    lact = int(row[4]) if len(row) > 4 and row[4].strip().isdigit() else None
                    w = float(row[5].replace(',', '.')) if len(row) > 5 and row[5].strip() else None
                    notes = row[6].strip() if len(row) > 6 and row[6].strip() else None
                    
                    if cat not in ['падёж', 'выбраковка', 'санитарный']:
                        errors.append(f"Строка {line_num}: недопустимая категория '{cat}'")
                        continue
                    if d_date > today:
                        errors.append(f"Строка {line_num}: дата {d_date} позже сегодняшней")
                        continue
                    rows_to_insert.append((cat, rea, d_date, ear_tag, lact, w, notes))
                    
            elif filename.endswith(('.xlsx', '.xls')):
                wb = openpyxl.load_workbook(file, data_only=True)
                ws = wb.active
                for line_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                    if not row or not row[0]:
                        continue
                    cat = str(row[0]).strip().lower()
                    rea = str(row[1]).strip() if row[1] else ''
                    
                    raw_date = row[2]
                    if isinstance(raw_date, datetime):
                        d_date = raw_date.strftime('%Y-%m-%d')
                    else:
                        d_date = str(raw_date).strip()
                        
                    ear_tag = str(row[3]).strip() if len(row) > 3 and row[3] else None
                    lact = int(row[4]) if len(row) > 4 and str(row[4]).strip().isdigit() else None
                    try:
                        w = float(str(row[5]).replace(',', '.')) if len(row) > 5 and row[5] else None
                    except ValueError:
                        w = None
                    notes = str(row[6]).strip() if len(row) > 6 and row[6] else None
                    
                    if cat not in ['падёж', 'выбраковка', 'санитарный']:
                        errors.append(f"Строка {line_num}: недопустимая категория '{cat}'")
                        continue
                    if d_date > today:
                        errors.append(f"Строка {line_num}: дата {d_date} позже сегодняшней")
                        continue
                    rows_to_insert.append((cat, rea, d_date, ear_tag, lact, w, notes))
            else:
                flash('Поддерживаются только форматы .xlsx, .xls и .csv', 'danger')
                return redirect(request.url)
                
        except Exception as e:
            flash(f'Ошибка обработки файла: {str(e)}', 'danger')
            return redirect(request.url)
            
        if rows_to_insert:
            conn = get_db_connection()
            cursor = conn.cursor()
            inserted_count = 0
            for r in rows_to_insert:
                cow_id = generate_cow_id(farm_name)
                cursor.execute('''
                    INSERT INTO cows (cow_id, ear_tag, farm_name, category, reason, disposal_date, lactation, weight, notes, created_by)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                ''', (cow_id, r[3], farm_name, r[0], r[1], r[2], r[4], r[5], r[6], session.get('user_id')))
                inserted_count += 1
            conn.commit()
            conn.close()
            
            log_audit_event(
                username=session.get('username'),
                farm_name=farm_name,
                action='IMPORT',
                entity_type='COW',
                details=f"Импортировано {inserted_count} записей из файла {file.filename}",
                user_id=session.get('user_id'),
                ip_address=request.remote_addr
            )
            
            flash(f'Успешно импортировано {inserted_count} коров!', 'success')
            if errors:
                flash(f'Внимание: пропущено строк с ошибками: {len(errors)}', 'warning')
            return redirect(url_for('cows.view_cows'))
        else:
            flash('В файле не найдено корректных строк для импорта.', 'danger')
            return redirect(request.url)
            
    return render_template('import_cows.html', farm_name=farm_name)

@cows_bp.route('/download_import_template')
@login_required
def download_import_template():
    """Скачать образец Excel-шаблона для импорта"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Шаблон_Импорта"
    
    headers = ['Категория (падёж/выбраковка/санитарный)*', 'Причина выбытия*', 'Дата (ГГГГ-ММ-ДД)*', 'Инв./Ушной №', 'Лактация', 'Масса (кг)', 'Примечание']
    ws.append(headers)
    
    # Примерные строки
    ws.append(['падёж', 'Острая тимпания', datetime.now().strftime('%Y-%m-%d'), '7701', 3, 520.5, 'Заключение ветврача'])
    ws.append(['выбраковка', 'Агалактия', datetime.now().strftime('%Y-%m-%d'), '7702', 4, 480.0, 'Низкая продуктивность'])
    ws.append(['санитарный', 'Травма конечностей', datetime.now().strftime('%Y-%m-%d'), '7703', 2, 510.0, 'Травма в секции'])
    
    # Стилизация шапки
    for cell in ws[1]:
        cell.font = openpyxl.styles.Font(bold=True, color="FFFFFF")
        cell.fill = openpyxl.styles.PatternFill(start_color="3B82F6", end_color="3B82F6", fill_type="solid")
        cell.alignment = openpyxl.styles.Alignment(horizontal="center")
        
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 14)
        
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return send_file(out, download_name="shablon_importa_korov.xlsx", as_attachment=True, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
